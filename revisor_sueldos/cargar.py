"""Lectura de planillas de remuneraciones en formato Transparencia Activa (CSV o Excel).

Cada fila queda como un dict normalizado:
    periodo ("AAAA-MM"), anio, mes, tipo (planta/contrata/honorarios/codigo_trabajo),
    estamento (texto original), grupo (académicos, directivos, ...), persona (clave),
    nombre, grado (int o None), jornada (horas o None), jerarquia (texto o None), rango_jerarquia (int o None),
    cargo, bruto (float), liquido (float o None), archivo.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from pathlib import Path

MESES = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6, "julio": 7,
    "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12,
    "ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6, "jul": 7, "ago": 8,
    "sep": 9, "sept": 9, "oct": 10, "nov": 11, "dic": 12,
}

# nombre lógico -> fragmentos (sin tildes, minúsculas) que lo identifican en el encabezado
COLUMNAS = {
    "anio": ["ano", "anio", "año"],
    "mes": ["mes"],
    "estamento": ["estamento", "planta o escalafon", "escalafon"],
    "apellido_paterno": ["apellido paterno"],
    "apellido_materno": ["apellido materno"],
    "nombres": ["nombres", "nombre"],
    "grado": ["grado eus", "grado", "jornada"],
    "calificacion": ["calificacion profesional", "formacion"],
    "cargo": ["cargo o funcion", "cargo", "funcion"],
    "bruto": ["remuneracion bruta", "bruta mensualizada", "honorario total bruto", "renta bruta", "bruto"],
    "liquido": ["remuneracion liquida", "liquida mensualizada", "liquido"],
    "tipo": ["tipo de contrato", "tipo contrato", "calidad juridica", "tipo"],
}

TIPOS = [
    ("honorarios", r"honorario"),
    ("codigo_trabajo", r"codigo[\s_-]*(del?[\s_-]*)?trabajo"),
    ("contrata", r"contrata"),
    ("planta", r"planta"),
]


def sin_tildes(texto) -> str:
    texto = unicodedata.normalize("NFKD", str(texto or ""))
    return "".join(c for c in texto if not unicodedata.combining(c)).lower().strip()


def monto(valor) -> float | None:
    """'$ 1.234.567' -> 1234567.0 ; '1.234,5' -> 1234.5 ; 1234567 -> 1234567.0"""
    if valor is None:
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    s = re.sub(r"[^\d,.\-]", "", str(valor))
    if not s or s in "-.,":
        return None
    if "," in s:  # formato chileno con decimales
        s = s.replace(".", "").replace(",", ".")
    elif s.count(".") > 1 or re.search(r"\.\d{3}$", s):
        s = s.replace(".", "")  # puntos de miles
    try:
        return float(s)
    except ValueError:
        return None


def clasificar_grupo(estamento: str, cargo: str, reglas: dict[str, list[str]]) -> str:
    """Primera regla (en orden) cuyo patrón calce con el estamento o, en su defecto, con el cargo."""
    for texto in (sin_tildes(estamento), sin_tildes(cargo)):
        if not texto:
            continue
        for grupo, patrones in reglas.items():
            if any(re.search(p, texto) for p in patrones):
                return grupo
    return "sin_clasificar"


def jerarquia(texto: str, rangos: dict[str, int]) -> tuple[str | None, int | None]:
    t = sin_tildes(texto)
    mejor = (None, None)
    for nombre, rango in rangos.items():
        if re.search(rf"\b{sin_tildes(nombre)}", t) and (mejor[1] is None or rango > mejor[1]):
            mejor = (nombre, rango)
    return mejor


def grado(valor) -> int | None:
    if isinstance(valor, (int, float)):
        return int(valor)
    s = sin_tildes(valor)
    if "hora" in s:  # "44 horas" es jornada, no grado
        return None
    m = re.search(r"\d+", s)
    return int(m.group()) if m else None


def jornada(valor) -> int | None:
    """'44 horas' -> 44. Sólo cuando el texto habla de horas."""
    s = sin_tildes(valor)
    m = re.search(r"(\d+)\s*(hrs?|horas?)", s)
    return int(m.group(1)) if m else None


def periodo_desde_texto(texto: str) -> tuple[int | None, int | None]:
    t = sin_tildes(texto)
    anio = re.search(r"(20\d\d)", t)
    mes = None
    for nombre, num in sorted(MESES.items(), key=lambda kv: -len(kv[0])):
        if re.search(rf"(?<![a-z]){nombre}(?![a-z])", t):
            mes = num
            break
    if mes is None:
        m = re.search(r"20\d\d[-_ .](\d{1,2})(?!\d)", t)
        if m and 1 <= int(m.group(1)) <= 12:
            mes = int(m.group(1))
    return (int(anio.group(1)) if anio else None), mes


def tipo_desde_texto(texto: str) -> str | None:
    t = sin_tildes(texto)
    for tipo, patron in TIPOS:
        if re.search(patron, t):
            return tipo
    return None


def _mapear_encabezado(encabezado: list) -> dict[str, int]:
    norm = [sin_tildes(h) for h in encabezado]
    usados, mapa = set(), {}
    for logico, claves in COLUMNAS.items():
        for clave in claves:  # las claves van de más a menos específicas
            for i, h in enumerate(norm):
                if i in usados:
                    continue
                if (h == clave) or (len(clave) > 4 and clave in h):
                    mapa[logico] = i
                    usados.add(i)
                    break
            if logico in mapa:
                break
    return mapa


def _hojas(ruta: Path):
    """Entrega (nombre_hoja, filas) del archivo."""
    if ruta.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        libro = load_workbook(ruta, read_only=True, data_only=True)
        for hoja in libro.worksheets:
            yield hoja.title, [list(f) for f in hoja.iter_rows(values_only=True)]
        libro.close()
        return
    crudo = ruta.read_bytes()
    for codif in ("utf-8-sig", "latin-1"):
        try:
            texto = crudo.decode(codif)
            break
        except UnicodeDecodeError:
            continue
    muestra = texto[:5000]
    try:
        dialecto = csv.Sniffer().sniff(muestra, delimiters=";,\t|")
    except csv.Error:
        dialecto = csv.excel
        dialecto.delimiter = ";" if muestra.count(";") > muestra.count(",") else ","
    yield "", list(csv.reader(io.StringIO(texto), dialecto))


def _fila_encabezado(filas: list[list]) -> int | None:
    """Algunas planillas traen títulos antes del encabezado: se busca la fila con más columnas reconocidas."""
    mejor, idx = 0, None
    for i, fila in enumerate(filas[:30]):
        mapa = _mapear_encabezado(fila)
        puntaje = len(mapa) + (3 if "bruto" in mapa else 0)
        if puntaje > mejor:
            mejor, idx = puntaje, i
    return idx if mejor >= 5 else None


def leer_archivo(ruta: Path, config: dict) -> list[dict]:
    reglas = config["grupos"]
    rangos = config.get("jerarquias_academicas", {})
    registros = []
    for nombre_hoja, filas in _hojas(ruta):
        i = _fila_encabezado(filas)
        if i is None:
            continue
        mapa = _mapear_encabezado(filas[i])
        anio_arch, mes_arch = periodo_desde_texto(f"{ruta.stem} {nombre_hoja}")
        tipo_arch = tipo_desde_texto(f"{ruta.stem} {nombre_hoja}")
        val = lambda fila, c: fila[mapa[c]] if c in mapa and mapa[c] < len(fila) else None  # noqa: E731
        for fila in filas[i + 1:]:
            if not fila or all(v in (None, "") for v in fila):
                continue
            bruto = monto(val(fila, "bruto"))
            if bruto is None:
                continue
            anio = grado(val(fila, "anio")) or anio_arch
            mes_v = val(fila, "mes")
            mes = (int(mes_v) if isinstance(mes_v, (int, float)) else periodo_desde_texto(f"2000 {mes_v}")[1]
                   if mes_v not in (None, "") else None) or mes_arch
            if not anio or not mes:
                raise ValueError(f"{ruta.name}: no se pudo determinar año/mes (agrega columnas Año/Mes "
                                 "o pon el período en el nombre del archivo, p. ej. 'planta_2025_03.csv')")
            partes = [val(fila, c) for c in ("apellido_paterno", "apellido_materno", "nombres")]
            nombre = " ".join(str(p).strip() for p in partes if p not in (None, ""))
            estamento = str(val(fila, "estamento") or "")
            cargo = " ".join(str(val(fila, c) or "") for c in ("cargo", "calificacion")).strip()
            jer, rango = jerarquia(f"{estamento} {cargo}", rangos)
            tipo = tipo_desde_texto(str(val(fila, "tipo") or "")) or tipo_arch or "sin_tipo"
            registros.append({
                "periodo": f"{anio:04d}-{mes:02d}", "anio": anio, "mes": mes, "tipo": tipo,
                "estamento": estamento.strip(), "grupo": clasificar_grupo(estamento, cargo, reglas),
                "persona": re.sub(r"\s+", " ", sin_tildes(nombre)).upper(), "nombre": nombre,
                "grado": grado(val(fila, "grado")), "jornada": jornada(val(fila, "grado")), "jerarquia": jer, "rango_jerarquia": rango,
                "cargo": cargo, "bruto": bruto, "liquido": monto(val(fila, "liquido")),
                "archivo": ruta.name,
            })
    return registros


def cargar(rutas: list[Path], config: dict) -> list[dict]:
    archivos = []
    for r in rutas:
        r = Path(r)
        if r.is_dir():
            archivos += sorted(p for p in r.rglob("*") if p.suffix.lower() in (".csv", ".xlsx", ".xlsm", ".txt"))
        else:
            archivos.append(r)
    registros = []
    for a in archivos:
        registros += leer_archivo(a, config)
    if not registros:
        raise ValueError("No se encontraron filas con remuneración bruta en las planillas entregadas")
    return registros
