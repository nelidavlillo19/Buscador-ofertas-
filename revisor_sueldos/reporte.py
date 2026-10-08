"""Salidas: libro Excel con todas las tablas e informe HTML autocontenido."""
from __future__ import annotations

import html
from pathlib import Path

NOMBRES_GRUPO = {"academicos": "Académicos", "directivos": "Directivos", "profesionales": "Profesionales",
                 "tecnicos": "Técnicos", "administrativos": "Administrativos", "auxiliares": "Auxiliares",
                 "sin_clasificar": "Sin clasificar", "TOTAL": "TOTAL"}
PORCENTAJES = ("tasa", "variacion", "diferencia_pct", "premio", "aumento", "incentivo_medio")


def _es_pct(col: str) -> bool:
    return any(col.startswith(p) or p in col for p in PORCENTAJES)


def pesos(v) -> str:
    return "—" if v is None else f"${v:,.0f}".replace(",", ".")


def millones(v) -> str:
    return "—" if v is None else f"${v / 1e6:,.1f} MM".replace(",", "X").replace(".", ",").replace("X", ".")


def pct(v) -> str:
    return "—" if v is None else f"{v * 100:.1f}%".replace(".", ",")


# ---------------------------------------------------------------- Excel

def excel(resultado: dict, ruta: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    hojas = [
        ("Costeo anual", resultado["anual"]),
        ("Costeo mensual", resultado["periodo"]),
        ("Costeo por contrato", resultado["tipo"]),
        ("Tasas de ascenso", resultado["trayectorias"]["tasas"]),
        ("Movimientos", resultado["trayectorias"]["movimientos"]),
        ("Proyección", resultado["proyeccion"]["filas"]),
        ("Proyección niveles", resultado["proyeccion"]["niveles"]),
        ("Supuestos proyección", resultado["proyeccion"]["supuestos"]),
        ("Encasillamiento", resultado["proyeccion"]["encasillamiento"]),
        ("Hallazgos revisión", resultado["hallazgos"]),
    ]
    libro = Workbook()
    libro.remove(libro.active)
    for titulo, filas in hojas:
        hoja = libro.create_sheet(titulo[:31])
        if not filas:
            hoja.append(["Sin datos"])
            continue
        columnas = list(filas[0].keys())
        hoja.append(columnas)
        for f in filas:
            hoja.append([f.get(c) for c in columnas])
        for i, c in enumerate(columnas, start=1):
            celda = hoja.cell(row=1, column=i)
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="1C5CAB")
            formato = "0.0%" if _es_pct(c) else ("#,##0" if isinstance(filas[0].get(c), float) else None)
            if formato:
                for fila in hoja.iter_rows(min_row=2, min_col=i, max_col=i):
                    fila[0].number_format = formato
            hoja.column_dimensions[get_column_letter(i)].width = max(12, min(len(c) + 4, 40))
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions
    libro.save(ruta)


# ---------------------------------------------------------------- HTML

def _tabla(filas: list[dict], columnas: list[tuple[str, str, str]], limite: int | None = None) -> str:
    """columnas: (clave, título, formato) con formato en {'pesos','mm','pct','num','txt','grupo'}"""
    fmt = {"pesos": pesos, "mm": millones, "pct": pct, "txt": lambda v: "—" if v is None else str(v),
           "num": lambda v: "—" if v is None else f"{v:,.1f}".replace(",", "X").replace(".", ",").replace("X", ".")
           if isinstance(v, float) and v % 1 else f"{v:,.0f}".replace(",", "."),
           "grupo": lambda v: NOMBRES_GRUPO.get(v, v)}
    cab = "".join(f"<th class='{'n' if f not in ('txt', 'grupo') else ''}'>{html.escape(t)}</th>"
                  for _, t, f in columnas)
    cuerpo = []
    for f in filas[:limite] if limite else filas:
        clase = " class='total'" if f.get("grupo") == "TOTAL" else ""
        celdas = "".join(f"<td class='{'n' if fm not in ('txt', 'grupo') else ''}'>{html.escape(fmt[fm](f.get(k)))}</td>"
                         for k, _, fm in columnas)
        cuerpo.append(f"<tr{clase}>{celdas}</tr>")
    extra = (f"<p class='nota'>Mostrando {limite} de {len(filas)} filas; el detalle completo está en el Excel.</p>"
             if limite and len(filas) > limite else "")
    return f"<div class='tabla'><table><thead><tr>{cab}</tr></thead><tbody>{''.join(cuerpo)}</tbody></table></div>{extra}"


def _grafico_proyeccion(filas: list[dict]) -> str:
    tot = [f for f in filas if f["grupo"] == "TOTAL"]
    if not tot:
        return ""
    ancho, alto, iz, de, ar, ab = 680, 280, 100, 130, 16, 32
    maximo = max(max(f["costo_base"], f["costo_reforma_total"]) for f in tot) * 1.05
    minimo = min(min(f["costo_base"], f["costo_reforma_total"]) for f in tot) * 0.9  # eje sin cero: importa la brecha
    x = lambda i: iz + (ancho - iz - de) * (i / max(len(tot) - 1, 1))  # noqa: E731
    y = lambda v: ar + (alto - ar - ab) * (1 - (v - minimo) / (maximo - minimo))  # noqa: E731
    partes = []
    for k in range(5):
        v = minimo + (maximo - minimo) * k / 4
        partes.append(f"<line x1='{iz}' x2='{ancho - de}' y1='{y(v):.1f}' y2='{y(v):.1f}' class='grid'/>"
                      f"<text x='{iz - 8}' y='{y(v) + 4:.1f}' class='eje' text-anchor='end'>{millones(v)}</text>")
    for i, f in enumerate(tot):
        partes.append(f"<text x='{x(i):.1f}' y='{alto - 10}' class='eje' text-anchor='middle'>{f['anio']}</text>")
    for clave, clase, etiqueta in (("costo_base", "s1", "Sin reforma"), ("costo_reforma_total", "s2", "Con nuevo sistema")):
        pts = " ".join(f"{x(i):.1f},{y(f[clave]):.1f}" for i, f in enumerate(tot))
        partes.append(f"<polyline points='{pts}' class='linea {clase}'/>")
        for i, f in enumerate(tot):
            tip = html.escape(f"{f['anio']} · {etiqueta}: {millones(f[clave])}")
            partes.append(f"<circle cx='{x(i):.1f}' cy='{y(f[clave]):.1f}' r='4' class='punto {clase}'/>"
                          f"<circle cx='{x(i):.1f}' cy='{y(f[clave]):.1f}' r='12' class='hit' data-tip='{tip}'/>")
        u = tot[-1]
        partes.append(f"<text x='{x(len(tot) - 1) + 10:.1f}' y='{y(u[clave]) + 4:.1f}' class='rotulo'>{etiqueta}</text>")
    return (f"<svg viewBox='0 0 {ancho} {alto}' class='grafico' role='img' "
            f"aria-label='Costo anual total proyectado con y sin el nuevo sistema'>{''.join(partes)}</svg>")


def _grafico_barras(datos: list[tuple[str, float]], formato) -> str:
    if not datos:
        return ""
    maximo = max(v for _, v in datos) or 1
    filas = []
    for etiqueta, v in datos:
        filas.append(f"<div class='barra' data-tip='{html.escape(f'{etiqueta}: {formato(v)}')}'>"
                     f"<span class='etq'>{html.escape(etiqueta)}</span>"
                     f"<span class='pista'><span class='relleno' style='width:{v / maximo * 100:.1f}%'></span></span>"
                     f"<span class='val'>{formato(v)}</span></div>")
    return f"<div class='barras'>{''.join(filas)}</div>"


def informe_html(resultado: dict, ruta: Path) -> None:
    anual = resultado["anual"]
    tray = resultado["trayectorias"]
    proy = resultado["proyeccion"]
    ultimo_anio = max(a["anio"] for a in anual)
    total_ult = next(a for a in anual if a["anio"] == ultimo_anio and a["grupo"] == "TOTAL")
    tot_proy = [f for f in proy["filas"] if f["grupo"] == "TOTAL"]
    fin = tot_proy[-1]
    acumulado = sum(f["diferencia"] for f in tot_proy)
    periodos = sorted({p["periodo"] for p in resultado["periodo"]})
    kpis = [
        ("Períodos revisados", f"{len(periodos)}", f"{periodos[0]} a {periodos[-1]}"),
        (f"Costo bruto anual {ultimo_anio}", millones(total_ult["bruto_anualizado"]),
         f"{total_ult['meses_con_datos']} meses con datos" + (" (anualizado)" if total_ult["meses_con_datos"] < 12 else "")),
        ("Dotación promedio", f"{total_ult['dotacion_promedio']:,.0f}".replace(",", "."), "personas por mes"),
        (f"Costo adicional {fin['anio']}", millones(fin["diferencia"]), f"{pct(fin['diferencia_pct'])} sobre escenario sin reforma"),
        ("Costo adicional acumulado", millones(acumulado), f"{tot_proy[1]['anio']}–{fin['anio']}"),
    ]
    kpi_html = "".join(f"<div class='kpi'><div class='k-t'>{html.escape(t)}</div><div class='k-v'>{html.escape(v)}</div>"
                       f"<div class='k-s'>{html.escape(s)}</div></div>" for t, v, s in kpis)
    asc = sorted(((NOMBRES_GRUPO.get(g, g), r["tasa_ascenso_anual"]) for g, r in tray["resumen"].items()),
                 key=lambda x: -x[1])
    costo_grupos = sorted(((NOMBRES_GRUPO.get(a["grupo"], a["grupo"]), a["bruto_anualizado"]) for a in anual
                           if a["anio"] == ultimo_anio and a["grupo"] != "TOTAL"), key=lambda x: -x[1])
    obs_tipos = {}
    for h in resultado["hallazgos"]:
        clave = h["observacion"].split("(")[0].split(":")[0].strip()
        obs_tipos[clave] = obs_tipos.get(clave, 0) + 1

    secciones = f"""
<section><h2>1. Costeo por grupo de trabajadores</h2>
<p>Remuneración bruta mensualizada informada en las planillas. El costo para el empleador agrega los aportes
configurados por tipo de contrato. Años incompletos se anualizan con el promedio de los meses disponibles.</p>
<h3>Costo bruto anual {ultimo_anio} por grupo</h3>
{_grafico_barras(costo_grupos, millones)}
{_tabla(anual, [("anio", "Año", "txt"), ("grupo", "Grupo", "grupo"), ("meses_con_datos", "Meses", "num"),
                ("dotacion_promedio", "Dotación prom.", "num"), ("bruto_anualizado", "Bruto anual", "mm"),
                ("costo_empleador_anualizado", "Costo empleador", "mm"),
                ("costo_medio_por_persona_anual", "Costo medio/persona", "pesos"), ("variacion_anual", "Var. anual", "pct")])}
<h3>Último período ({periodos[-1]})</h3>
{_tabla([p for p in resultado["periodo"] if p["periodo"] == periodos[-1]],
        [("grupo", "Grupo", "grupo"), ("dotacion", "Personas", "num"), ("total_bruto", "Total bruto", "mm"),
         ("promedio", "Promedio", "pesos"), ("mediana", "Mediana", "pesos"), ("p10", "P10", "pesos"),
         ("p90", "P90", "pesos"), ("maximo", "Máximo", "pesos")])}
</section>
<section><h2>2. Revisión hacia atrás: ascensos y movilidad</h2>
<p>Se compara a cada persona entre los cortes {", ".join(tray["cortes"])}. Ascenso = sube de grado EUS,
sube de jerarquía académica o pasa a un estamento superior. La tasa anual corrige por la distancia en meses.
El premio de ascenso es el aumento mediano de los ascendidos menos el del resto.</p>
<h3>Tasa anual de ascenso por grupo (promedio ponderado)</h3>
{_grafico_barras(asc, pct)}
{_tabla(tray["tasas"], [("desde", "Desde", "txt"), ("hasta", "Hasta", "txt"), ("grupo", "Grupo", "grupo"),
                        ("dotacion_inicial", "Dot. inicial", "num"), ("permanecen", "Permanecen", "num"),
                        ("ingresos", "Ingresos", "num"), ("egresos", "Egresos", "num"),
                        ("ascendidos", "Ascendidos", "num"), ("tasa_ascenso_anual", "Tasa ascenso anual", "pct"),
                        ("grados_promedio_por_ascenso", "Grados/ascenso", "num"),
                        ("premio_ascenso", "Premio ascenso", "pct"), ("pasos_a_planta", "A planta", "num"),
                        ("tasa_egreso_anual", "Egreso anual", "pct")])}
<h3>Movimientos individuales</h3>
{_tabla(tray["movimientos"], [("hasta", "Hasta", "txt"), ("persona", "Persona", "txt"),
                              ("grupo_inicial", "Grupo", "grupo"), ("movimiento", "Movimiento", "txt"),
                              ("grado_inicial", "Grado ini.", "txt"), ("grado_final", "Grado fin.", "txt"),
                              ("jerarquia_final", "Jerarquía", "txt"), ("variacion_bruto", "Var. bruto", "pct")], 40)}
</section>
<section><h2>3. Proyección del nuevo sistema de jerarquización y evaluación</h2>
<p>Base: {proy["periodo_base"]}. Brecha mensual inicial para llevar a todos al piso de su nivel:
<strong>{pesos(proy["brecha_total_mensual"])}</strong> ({millones(proy["brecha_total_mensual"] * 12)} al año a régimen,
antes de reajustes). Ver supuestos al final.</p>
<h3>Costo anual total: con y sin nuevo sistema</h3>
<div class='leyenda'><span><i class='m s1'></i>Sin reforma</span><span><i class='m s2'></i>Con nuevo sistema (incluye incentivo por evaluación)</span></div>
{_grafico_proyeccion(proy["filas"])}
{_tabla(proy["filas"], [("anio", "Año", "txt"), ("grupo", "Grupo", "grupo"), ("dotacion", "Dotación", "num"),
                        ("costo_base", "Sin reforma", "mm"), ("costo_reforma_remuneraciones", "Remuneraciones", "mm"),
                        ("incentivo_evaluacion", "Incentivo eval.", "mm"), ("costo_reforma_total", "Con reforma", "mm"),
                        ("diferencia", "Diferencia", "mm"), ("diferencia_pct", "Dif. %", "pct"),
                        ("costo_empleador_reforma", "Costo empleador", "mm")])}
<h3>Supuestos por grupo</h3>
{_tabla(proy["supuestos"], [("grupo", "Grupo", "grupo"), ("personas", "Personas", "num"),
                            ("tasa_ascenso_historica", "Tasa ascenso", "pct"), ("origen_tasa", "Origen", "txt"),
                            ("tasa_ascenso_con_evaluacion", "Tasa c/evaluación", "pct"),
                            ("premio_ascenso", "Premio ascenso", "pct"), ("incentivo_medio", "Incentivo medio", "pct")])}
<h3>Distribución esperada por nivel (último año proyectado)</h3>
{_tabla([n for n in proy["niveles"] if n["anio"] == fin["anio"]],
        [("grupo", "Grupo", "grupo"), ("nivel", "Nivel", "txt"), ("personas_esperadas", "Personas", "num"),
         ("remuneracion_media", "Remuneración media", "pesos")])}
</section>
<section><h2>4. Hallazgos de la revisión</h2>
<p>{len(resultado["hallazgos"])} observaciones. Resumen por tipo:</p>
{_tabla([{"tipo": k, "n": v} for k, v in sorted(obs_tipos.items(), key=lambda x: -x[1])],
        [("tipo", "Observación", "txt"), ("n", "Casos", "num")])}
{_tabla(resultado["hallazgos"], [("periodo", "Período", "txt"), ("persona", "Persona", "txt"),
                                 ("grupo", "Grupo", "grupo"), ("bruto", "Bruto", "pesos"),
                                 ("observacion", "Observación", "txt")], 40)}
</section>"""
    ruta.write_text(PLANTILLA.replace("{{KPIS}}", kpi_html).replace("{{SECCIONES}}", secciones)
                    .replace("{{FUENTE}}", html.escape(resultado["fuente"])), encoding="utf-8")


PLANTILLA = """<!doctype html>
<html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Revisión de remuneraciones</title>
<style>
:root{color-scheme:light;--bg:#fcfcfb;--panel:#ffffff;--tx:#0b0b0b;--tx2:#52514e;--borde:#e4e3df;--grid:#ecebe7;
--s1:#2a78d6;--s2:#eb6834;--acento:#1c5cab;--total:#f3f2ee}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){color-scheme:dark;--bg:#1a1a19;--panel:#222221;
--tx:#ffffff;--tx2:#c3c2b7;--borde:#383835;--grid:#2e2e2c;--s1:#3987e5;--s2:#d95926;--acento:#86b6ef;--total:#2a2a28}}
:root[data-theme="dark"]{color-scheme:dark;--bg:#1a1a19;--panel:#222221;--tx:#ffffff;--tx2:#c3c2b7;--borde:#383835;
--grid:#2e2e2c;--s1:#3987e5;--s2:#d95926;--acento:#86b6ef;--total:#2a2a28}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--tx);font:15px/1.5 system-ui,-apple-system,"Segoe UI",sans-serif}
main{max-width:1100px;overflow-wrap:anywhere;margin:0 auto;padding:24px 16px 64px}h1{font-size:1.6rem;margin:0 0 4px}
h2{font-size:1.2rem;margin:40px 0 8px;padding-top:16px;border-top:1px solid var(--borde)}h3{font-size:1rem;margin:24px 0 8px}
p{color:var(--tx2);max-width:75ch}.sub{color:var(--tx2);margin:0 0 20px}
.kpis{display:grid;grid-template-columns:repeat(auto-fit,minmax(180px,1fr));gap:12px}
.kpi{min-width:0;background:var(--panel);border:1px solid var(--borde);border-radius:10px;padding:12px 14px}
.k-t{font-size:.8rem;color:var(--tx2)}.k-v{font-size:1.4rem;font-weight:650;font-variant-numeric:tabular-nums}.k-s{font-size:.78rem;color:var(--tx2)}
.tabla{overflow-x:auto;border:1px solid var(--borde);border-radius:8px;background:var(--panel);margin:8px 0}
table{border-collapse:collapse;width:100%;font-size:.85rem}th,td{padding:6px 10px;border-bottom:1px solid var(--grid);white-space:nowrap;text-align:left}
th{position:sticky;top:0;background:var(--panel);color:var(--tx2);font-weight:600}.n{text-align:right;font-variant-numeric:tabular-nums}
tr.total td{background:var(--total);font-weight:600}.nota{font-size:.8rem}
.barras{display:grid;gap:6px;margin:8px 0 16px;max-width:720px}.barra{display:grid;grid-template-columns:130px 1fr 110px;gap:10px;align-items:center;font-size:.85rem}
.pista{height:14px;min-width:0}.etq{overflow:hidden;text-overflow:ellipsis;white-space:nowrap}.relleno{display:block;height:100%;background:var(--s1);border-radius:0 4px 4px 0}.val{text-align:right;font-variant-numeric:tabular-nums}
.grafico{width:100%;max-width:760px;height:auto;display:block}.grid{stroke:var(--grid)}.eje{fill:var(--tx2);font-size:11px}
.rotulo{fill:var(--tx);font-size:12px}.linea{fill:none;stroke-width:2}.linea.s1{stroke:var(--s1)}.linea.s2{stroke:var(--s2)}
.punto{stroke:var(--panel);stroke-width:2}.punto.s1{fill:var(--s1)}.punto.s2{fill:var(--s2)}.hit{fill:transparent;cursor:pointer}
.leyenda{display:flex;gap:16px;flex-wrap:wrap;font-size:.85rem;color:var(--tx2)}.m{display:inline-block;width:14px;height:3px;border-radius:2px;margin-right:6px;vertical-align:middle}
.m.s1{background:var(--s1)}.m.s2{background:var(--s2)}
#tip{position:fixed;pointer-events:none;background:var(--panel);color:var(--tx);border:1px solid var(--borde);border-radius:6px;padding:4px 8px;font-size:.8rem;display:none;box-shadow:0 2px 8px rgba(0,0,0,.15)}
@media (max-width:600px){.barra{grid-template-columns:88px minmax(0,1fr) 84px;font-size:.78rem}.k-v{font-size:1.15rem}}
</style></head><body><main>
<h1>Revisión de remuneraciones y proyección de costos</h1>
<p class="sub">Fuente: {{FUENTE}}</p>
<div class="kpis">{{KPIS}}</div>
{{SECCIONES}}
<section><h2>Supuestos y limitaciones</h2><p>Las planillas de Transparencia no traen RUT: las personas se siguen por nombre
completo normalizado (homónimos o cambios de nombre pueden generar falsos ingresos/egresos). La proyección es en valor
esperado y a dotación constante; los parámetros (pisos de nivel, gradualidad, reajuste, distribución de la evaluación,
incentivos y factores de costo empleador) están en <code>config/revisor_sueldos.yaml</code> y deben validarse con
Finanzas y Recursos Humanos.</p></section>
</main><div id="tip"></div>
<script>
const tip=document.getElementById('tip');
document.addEventListener('mouseover',e=>{const t=e.target.closest('[data-tip]');if(!t){tip.style.display='none';return}
tip.textContent=t.dataset.tip;tip.style.display='block'});
document.addEventListener('mousemove',e=>{tip.style.left=(e.clientX+12)+'px';tip.style.top=(e.clientY+12)+'px'});
</script></body></html>"""
