from pathlib import Path

import pytest
import yaml

from revisor_sueldos import __main__ as cli
from revisor_sueldos import cargar, costeo, proyeccion, trayectorias

CONFIG = yaml.safe_load((Path(__file__).resolve().parent.parent / "config" / "revisor_sueldos.yaml")
                        .read_text(encoding="utf-8"))


def test_montos_en_formato_chileno():
    assert cargar.monto("$ 1.234.567") == 1234567
    assert cargar.monto("1.234,5") == 1234.5
    assert cargar.monto(980000) == 980000
    assert cargar.monto("") is None


def test_profesor_es_academico_y_no_profesional():
    reglas = CONFIG["grupos"]
    assert cargar.clasificar_grupo("Profesor", "", reglas) == "academicos"
    assert cargar.clasificar_grupo("Profesional", "", reglas) == "profesionales"
    assert cargar.clasificar_grupo("Técnico", "", reglas) == "tecnicos"
    assert cargar.clasificar_grupo("", "Profesor Asociado", reglas) == "academicos"


def test_periodo_y_tipo_desde_nombre_de_archivo():
    assert cargar.periodo_desde_texto("planta_marzo_2025") == (2025, 3)
    assert cargar.periodo_desde_texto("contrata 2026-07") == (2026, 7)
    assert cargar.tipo_desde_texto("Personal_Codigo_del_Trabajo") == "codigo_trabajo"
    assert cargar.jornada("44 horas") == 44 and cargar.grado("44 horas") is None


def _csv(ruta: Path, filas: list[str]) -> Path:
    ruta.write_text("\n".join(filas), encoding="utf-8")
    return ruta


ENC = "Estamento;Apellido paterno;Apellido materno;Nombres;Grado EUS o jornada;Cargo o función;Remuneración bruta mensualizada"


def test_lee_csv_con_titulos_previos_y_periodo_en_nombre(tmp_path):
    ruta = _csv(tmp_path / "planta_diciembre_2025.csv", [
        "Universidad;;;;;;", "Personal de planta;;;;;;", ENC,
        "Profesional;Soto;Díaz;Ana;12;Analista;1.800.000",
        "Académico;Rojas;Muñoz;Luis;44 horas;Profesor Asociado;3.100.000",
    ])
    filas = cargar.leer_archivo(ruta, CONFIG)
    assert [f["periodo"] for f in filas] == ["2025-12", "2025-12"]
    assert filas[0]["grupo"] == "profesionales" and filas[0]["grado"] == 12 and filas[0]["tipo"] == "planta"
    assert filas[1]["jerarquia"] == "asociado" and filas[1]["jornada"] == 44


def test_tasa_de_ascenso_y_egresos(tmp_path):
    a = _csv(tmp_path / "planta_2024_12.csv", [ENC] + [f"Profesional;P{i};X;Y;12;A;1.000.000" for i in range(10)])
    b = _csv(tmp_path / "planta_2025_12.csv", [ENC]
             + [f"Profesional;P{i};X;Y;{11 if i < 2 else 12};A;1.100.000" for i in range(9)])
    registros = cargar.cargar([a, b], CONFIG)
    t = trayectorias.analizar(registros, CONFIG)
    fila = t["tasas"][0]
    assert (fila["permanecen"], fila["ascendidos"], fila["egresos"]) == (9, 2, 1)
    assert fila["tasa_ascenso_anual"] == pytest.approx(2 / 9)
    assert len(t["movimientos"]) == 2


def test_costeo_anualiza_meses_faltantes(tmp_path):
    a = _csv(tmp_path / "planta_2025_01.csv", [ENC, "Auxiliar;A;B;C;25;Aux;600.000"])
    anual = costeo.anual(cargar.cargar([a], CONFIG), CONFIG)
    total = next(x for x in anual if x["grupo"] == "TOTAL")
    assert total["meses_con_datos"] == 1 and total["bruto_anualizado"] == 600000 * 12


def test_proyeccion_lleva_al_piso_gradualmente_y_respeta_garantia(tmp_path):
    a = _csv(tmp_path / "planta_2025_12.csv", [ENC, "Auxiliar;Bajo;B;C;25;Aux;500.000",
                                               "Auxiliar;Alto;B;C;25;Aux;900.000"])
    cfg = {**CONFIG, "proyeccion": {**CONFIG["proyeccion"], "reajuste_anual": 0.0, "tasas_ascenso": {"auxiliares": 0.0},
                                    "evaluacion": {}, "gradualidad": [0.5, 1.0], "horizonte_anios": 2}}
    registros = cargar.cargar([a], cfg)
    res = proyeccion.proyectar(trayectorias.consolidar(registros), {}, cfg)
    aux = {f["anio"]: f for f in res["filas"] if f["grupo"] == "auxiliares"}
    # piso Auxiliar C = 700.000: el de 500.000 sube la mitad de la brecha el año 1 y todo el año 2; el de 900.000 se mantiene
    assert aux[2026]["costo_reforma_total"] == pytest.approx((600000 + 900000) * 12)
    assert aux[2027]["costo_reforma_total"] == pytest.approx((700000 + 900000) * 12)
    assert res["brecha_total_mensual"] == pytest.approx(200000)


def test_ejemplo_completo_genera_informe_y_excel(tmp_path):
    rutas = cli.ejemplo.generar(tmp_path / "planillas")
    resultado = cli.analizar(rutas, CONFIG, "prueba")
    cli.escribir(resultado, tmp_path / "salida")
    assert (tmp_path / "salida" / "informe.html").stat().st_size > 10000
    assert (tmp_path / "salida" / "revision_sueldos.xlsx").exists()
    assert {g for g in resultado["trayectorias"]["resumen"]} >= {"academicos", "profesionales", "auxiliares"}
    assert any("doble vínculo" in h["observacion"] for h in resultado["hallazgos"])


def test_planilla_sin_apellidos_se_costea_pero_no_se_sigue(tmp_path):
    enc = "Año;Mes;Estamento;Primer nombre;Remuneración bruta mensual ($)"
    a = _csv(tmp_path / "planta_raw.csv", [enc, "2025;Enero;Auxiliar;Ana;600.000", "2025;Enero;Auxiliar;Ana;610.000"])
    registros = cargar.cargar([a], CONFIG)
    assert len({r["persona"] for r in registros}) == 2  # dos personas distintas aunque compartan nombre
    assert not costeo.hallazgos(registros, CONFIG)  # sin columna de grado no se reclama el grado
    assert trayectorias.analizar(registros, CONFIG)["tasas"] == []


def test_hojas_de_resumen_se_omiten(tmp_path):
    from openpyxl import Workbook
    libro = Workbook()
    resumen = libro.active
    resumen.title = "Hoja 01"
    resumen.append(["Año", "Mes", "Estamento", "Grado", "Rem. Bruta Planta ($)"])
    resumen.append(["TOTAL", "Todos", "", "", 999])
    nomina = libro.create_sheet("Contrata_Raw")
    nomina.append(["Período", "Estamento", "Apellido Paterno", "Nombres", "Grado", "Rem. Bruta ($)"])
    nomina.append(["2025-08", "Técnico", "Soto", "Ana", 18, 1000000])
    libro.create_sheet("Comisiones").append(["Período", "Nombres", "Apellido Paterno", "Gasto Bruto ($)"])
    libro.save(tmp_path / "libro.xlsx")
    filas = cargar.leer_archivo(tmp_path / "libro.xlsx", CONFIG)
    assert [(f["periodo"], f["tipo"], f["grupo"], f["bruto"]) for f in filas] == [("2025-08", "contrata", "tecnicos", 1000000)]


def test_anualiza_cada_tipo_con_sus_propios_meses(tmp_path):
    p1 = _csv(tmp_path / "planta_2025_01.csv", [ENC, "Auxiliar;A;B;C;25;Aux;600.000"])
    p2 = _csv(tmp_path / "planta_2025_02.csv", [ENC, "Auxiliar;A;B;C;25;Aux;600.000"])
    c1 = _csv(tmp_path / "contrata_2025_02.csv", [ENC, "Auxiliar;D;E;F;25;Aux;500.000"])
    total = next(x for x in costeo.anual(cargar.cargar([p1, p2, c1], CONFIG), CONFIG) if x["grupo"] == "TOTAL")
    # planta: 2 meses de 600.000; contrata: 1 mes de 500.000 → ambos a 12 meses
    assert total["bruto_anualizado"] == pytest.approx(600000 * 12 + 500000 * 12)
    assert total["meses_por_tipo"] == "contrata 1, planta 2"


def test_premio_por_nivel_y_tasa_distinta_con_reforma(tmp_path):
    a = _csv(tmp_path / "planta_2025_12.csv", [ENC, "Académico;A;B;C;44 horas;Profesor Asociado;2.000.000"])
    proy = {**CONFIG["proyeccion"], "reajuste_anual": 0.0, "evaluacion": {}, "gradualidad": [1.0], "horizonte_anios": 1,
            "tasas_ascenso": {"academicos": 0.0}, "tasas_ascenso_reforma": {"academicos": 0.5},
            "jerarquizacion": {"academicos": {"criterio": "jerarquia", "niveles": [
                {"nombre": "Asociado", "jerarquias": ["asociado"], "piso": 0},
                {"nombre": "Titular", "jerarquias": ["titular"], "piso": 0, "premio": 0.2}]}}}
    cfg = {**CONFIG, "proyeccion": proy}
    res = proyeccion.proyectar(trayectorias.consolidar(cargar.cargar([a], cfg)), {}, cfg)
    acad = {f["anio"]: f for f in res["filas"] if f["grupo"] == "academicos"}
    assert acad[2026]["costo_base"] == pytest.approx(2000000 * 12)  # sin reforma no asciende
    assert acad[2026]["costo_reforma_total"] == pytest.approx((0.5 * 2000000 + 0.5 * 2400000) * 12)
